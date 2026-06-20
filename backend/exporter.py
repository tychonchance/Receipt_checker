import io
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_YELLOW_FILL = PatternFill(start_color="FFF176", end_color="FFF176", fill_type="solid")
_RED_FILL = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")
_GREEN_FILL = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")
_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_HEADERS = [
    "Company",
    "Date",
    "Time",
    "Amount",
    "Currency",
    "Matched Event(s)",
    "Calendar(s)",
    "Status",
    "Notes",
]


def export_to_excel(receipts: List[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Receipts"

    for col, header in enumerate(_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER

    ws.row_dimensions[1].height = 20

    for row_num, receipt in enumerate(receipts, 2):
        status = receipt.get("status", "unmatched")
        matched_events = receipt.get("matched_events", [])

        fill = {
            "ambiguous": _YELLOW_FILL,
            "unmatched": _RED_FILL,
            "matched": _GREEN_FILL,
        }.get(status)

        event_names = ", ".join(e.get("summary", "") for e in matched_events)
        calendars = ", ".join(sorted({e.get("calendar_name", "") for e in matched_events}))
        amount = receipt.get("amount")

        row_data = [
            receipt.get("company") or "",
            receipt.get("date") or "",
            receipt.get("time") or "",
            f"{amount:.2f}" if amount is not None else "",
            receipt.get("currency") or "",
            event_names,
            calendars,
            status.capitalize(),
            receipt.get("raw_text") or "",
        ]

        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col, value=value)
            cell.border = _BORDER
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
            if fill:
                cell.fill = fill

    # Auto-fit column widths (capped at 50)
    for col in range(1, len(_HEADERS) + 1):
        max_len = len(_HEADERS[col - 1])
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col)].width = min(max_len + 3, 50)

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
